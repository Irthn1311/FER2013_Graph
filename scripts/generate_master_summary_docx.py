"""Script to generate 00_MASTER_SUMMARY_MPG_FER.docx using python-docx.

Produces a formatted, comprehensive scientific summary document for
audit paper preparation and Q1 experiment design.
"""

from __future__ import annotations

from pathlib import Path
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
STAGE_DIR = PROJECT_ROOT / "research" / "mpg_fer_handoff_stage"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = "Calibri"
        if level == 1:
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
            run.font.size = Pt(16)
            run.bold = True
        elif level == 2:
            run.font.color.rgb = RGBColor(0x2B, 0x5C, 0x8F)
            run.font.size = Pt(13)
            run.bold = True
        elif level == 3:
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            run.font.size = Pt(11)
            run.bold = True
    return h


def main():
    print("Building 00_MASTER_SUMMARY_MPG_FER.docx...", flush=True)

    doc = docx.Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("MPG-FER: Master Research & Architectural Summary")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Multi-Scale Pixel-Relational Motif Graph for Facial Expression Recognition on FER2013\nComprehensive Post-Hoc Audit Synthesis & Q1 Experimental Strategy")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Callout Box - Executive Status
    table_meta = doc.add_table(rows=1, cols=1)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_meta = table_meta.rows[0].cells[0]
    set_cell_background(cell_meta, "F0F4F8")
    set_cell_margins(cell_meta, 120, 120, 180, 180)
    p_meta = cell_meta.paragraphs[0]
    p_meta.paragraph_format.space_after = Pt(0)
    r_meta_b = p_meta.add_run("EXECUTIVE STATUS & SOURCE OF TRUTH:\n")
    r_meta_b.bold = True
    r_meta_b.font.size = Pt(10)
    r_meta_b.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    meta_text = (
        "• Repository: Irthn1311/FER2013_Graph (GitHub) | Branch: research/mpg-fer-v2-3-early-depth-generalization\n"
        "• Current Reference Version: v2.1 (Dense Complete) / v2.2 (Dynamic Top-K Sparse)\n"
        "• Current Best Verified Result: v2.3 (Private Raw Acc: 68.68%, Private TTA Acc: 70.66%, Macro-F1: 0.6998)\n"
        "• Audit Verdicts: A4 (Equivalence), A5-R (Leakage Closure), A5-H (Blinded AI Adjudication), A6-R2 (Source-Locked)\n"
        "• Recommended Q1 Target: EARLY_DEPTH_GENERALIZATION_TARGET (Layers 1–2 Residual Preservation)"
    )
    r_meta_t = p_meta.add_run(meta_text)
    r_meta_t.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 1. Architectural Philosophy & Pipeline
    # -------------------------------------------------------------
    add_styled_heading(doc, "1. Core Architectural Philosophy & Mathematical Pipeline", 1)
    p_intro = doc.add_paragraph()
    p_intro.add_run(
        "MPG-FER reformulates facial expression recognition as a hierarchical, two-level graph problem. "
        "Unlike conventional CNN or Vision Transformer backbones that rely on dense uniform attention across pixel patches, "
        "MPG-FER constructs: (1) a local edge-aware Pixel GNN over 2,304 pixel nodes, and (2) a high-level relational Motif Graph "
        "over 49 semantic facial tokens derived through multi-scale spatial prototype assignment."
    )

    p_pipe = doc.add_paragraph()
    p_pipe.add_run(
        "The complete forward pass is structured into five sequential computational stages:\n"
        "1. Pixel Descriptor & Topology: Extracts a 32-channel handcrafted descriptor (intensity, Sobel gradients, Hessian eigenvalues, Laplacian) "
        "and forms an 8-connected grid graph over 2,304 nodes with 5-channel continuous edge features.\n"
        "2. Edge-Aware Pixel GNN: A 4-layer spatial message-passing network with multi-head attention and residual drop-path updates pixel embeddings (96d). "
        "A multi-pooling projection produces a 128d global pixel readout vector.\n"
        "3. Multi-Scale Spatial Motif Composer: Evaluates 49 spatial anchor windows (stride 6, 7x7 grid) across three receptive field scales (8, 12, 16 pixels). "
        "For each anchor, it decomposes contextual pixel states into: WHAT (saliency-weighted visual appearance, 96d), TYPE (soft assignment over 48 learned prototype codebook vectors, 32d), "
        "and WHERE (continuous geometric moments cx, cy, sx, sy, mass, 5d). A learned scale gate fuses candidate scales into 49 occurrence tokens (192d).\n"
        "4. Geometry-Aware Motif Graph: A 5-layer transformer GNN that combines token self-attention with continuous 6d geometry bias (relative distances, angles, scales). "
        "In v2.2 and v2.3, attention is constrained to dynamic non-self Top-K supports [8, 16, 16, 16, 24].\n"
        "5. Dual Readout & Fusion Classifier: Multi-head attention pooling across the 49 motif tokens yields a 384d motif readout vector. "
        "Concatenation with the 128d pixel readout produces a 512d Fusion vector, classified by a 2-layer MLP (256d hidden, GELU, Dropout 0.25) into 7 discrete emotion classes."
    )

    # -------------------------------------------------------------
    # 2. Version Lineage & Benchmark Comparison
    # -------------------------------------------------------------
    add_styled_heading(doc, "2. Version Lineage & Comprehensive Benchmark Results", 1)
    p_lineage = doc.add_paragraph()
    p_lineage.add_run(
        "All official runs were executed on Kaggle Notebooks under identical training conditions (batch size 16, gradient accumulation 2, "
        "cosine LR decay, label smoothing 0.05, Model EMA decay 0.999). Parameter counts are strictly locked across v2.1, v2.2, and v2.3 at 2,304,528."
    )

    # Version comparison table
    tbl_ver = doc.add_table(rows=6, cols=9)
    tbl_ver.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_ver = ["Version", "Topology", "Top-K", "Res Scale", "Params", "Pub Raw", "Pub TTA", "Priv Raw", "Priv TTA (F1)"]
    for i, h in enumerate(headers_ver):
        cell = tbl_ver.rows[0].cells[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_ver = [
        ["v1", "Diffuse Proto", "48 (all)", "1.0 (all)", "2,219,788", "67.21%", "68.99%", "67.73%", "69.41% (0.687)"],
        ["v2", "Sharp Gumbel", "48 (all)", "1.0 (all)", "2,238,609", "66.98%", "68.46%", "67.96%", "69.43% (0.691)"],
        ["v2.1", "Dense Complete", "48 (all)", "1.0 (all)", "2,304,528", "67.43%", "69.27%", "68.21%", "69.94% (0.690)"],
        ["v2.2", "Dynamic Sparse", "8-16-16-16-24", "1.0 (all)", "2,304,528", "67.68%", "69.66%", "68.43%", "69.55% (0.687)"],
        ["v2.3", "Sparse + ResPreserve", "8-16-16-16-24", "0.5 (L1-2)", "2,304,528", "67.73%", "69.49%", "68.68%", "70.66% (0.700)"],
    ]
    for row_idx, row_vals in enumerate(data_ver):
        row = tbl_ver.rows[row_idx + 1]
        bg = "F9FAFC" if row_idx % 2 == 1 else "FFFFFF"
        if row_idx == 4:
            bg = "EBF3FA" # Highlight v2.3
        for col_idx, val in enumerate(row_vals):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if row_idx == 4:
                r.font.bold = True

    p_v23_note = doc.add_paragraph()
    p_v23_note.paragraph_format.space_before = Pt(6)
    p_v23_note.add_run(
        "Key Milestone: MPG-FER v2.3 breaks through the historical 70% threshold, reaching 70.66% Private TTA accuracy "
        "(Macro-F1: 0.6998) and a best-ever Private Raw accuracy of 68.68% (2,465 / 3,589). "
        "This gain was achieved with zero parameter increase, purely through early-depth residual preservation (res_scale=0.5 in Layers 1–2)."
    )

    # -------------------------------------------------------------
    # 3. Mechanistic Audit Synthesis (A1 to A6-R2)
    # -------------------------------------------------------------
    add_styled_heading(doc, "3. Complete Mechanistic Audit Synthesis", 1)
    p_audit_intro = doc.add_paragraph()
    p_audit_intro.add_run(
        "A cornerstone of the MPG-FER research program is experimental discipline: rather than stacking arbitrary architecture modifications, "
        "each design step is guided by frozen post-hoc mechanistic audits."
    )

    audits = [
        ("A1: Branch Bottleneck & Probe Audit",
         "Established that the relational Motif Graph is the primary driver of classification (reaching ~68% alone). "
         "The direct global Pixel readout is weak (~55%), and concatenating it at Fusion provides marginal linear gain (+0.8 pp). "
         "Linear probes on Composer subcomponents demonstrated that TYPE adds less than 1 pp over WHAT+WHERE."),
        
        ("A2 / A2-R: Graph Topology Falsification",
         "Demonstrated that full degree-48 connectivity is highly redundant. Pruned complete graphs to dynamic Top-K supports [8, 16, 16, 16, 24] "
         "without loss of validation accuracy. Established that dynamic sample-conditioned attention is strictly necessary: "
         "fixing static geometric 8-NN prior topologies caused immediate performance collapse (-1.5 pp)."),

        ("A4: Dense vs. Sparse Mechanistic Equivalence",
         "Direct comparison of frozen v2.1 and v2.2. Showed that macro-representations converge strongly (Fusion CKA = 0.857, Motif CKA = 0.837, strict diagonal cross-layer dominance). "
         "At the sample level, routing supports differ substantially (~50% overlap). Models share 73% of their errors. "
         "Linear probes underperformed the trained classifier (-1.5 pp), ruling out the classifier head as a bottleneck."),

        ("A5 / A5-R: Data Leakage & Annotation Inconsistency",
         "Cryptographic pixel SHA hashing uncovered 1,516 exact duplicate groups (3,369 images). Exactly 57 duplicate groups have conflicting ground-truth labels. "
         "Crucially, 288 PrivateTest rows (8.02%) have exact copies in Train, inflating benchmark scores by ~2.2 pp (leakage-excluded accuracy is ~67.4%). "
         "In unlearned raw pixel space, Fear exhibits an intrinsically negative centroid margin (-3.39), proving that hard-class overlap is an inherent property of the dataset."),

        ("A5-H: Blinded AI Visual Adjudication",
         "Two independent blinded AI visual reviewers (Reviewer C: Claude, Reviewer G: Gemini R2) evaluated 200 balanced test samples. "
         "Inter-rater agreement was 47.0% (Cohen's kappa 0.3811). On 50 high-confidence shared errors, strict blind consensus favored the model's prediction nearly 3x more often than the nominal label (34% vs. 12%, exact binomial p=0.035). "
         "Cataloged 17 strict Label Mismatch candidates and 11 Common Representation Failure candidates."),

        ("A6 / A6-R / A6-R2: Representation Failure Localization",
         "Tracked 1,197 model-resolvable samples across 11 stages and 9 swap boundaries (S0–S8). Reached exact source-lock closure:\n"
         "• S0 (Pixel Readout) rescues <2.2% (global readout bypass is inert).\n"
         "• S1 (PRE-Motif Composer) rescues ~36.9%, driven entirely by visual appearance (WHAT: 36.9%), while prototype assignment (TYPE) adds <0.5%.\n"
         "• Early Motif Reasoning (Layers 1–2) accelerates rescue from 36.4% to 57.8% (+21.4% gain) and doubles margin divergence.\n"
         "• S6 (Layer-5 Node States) alone rescues ~74% under receiver suffix, proving node states dominate classification.\n"
         "• S7 (Motif Readout) packages node states to reach ~89% rescue.\n"
         "• Generalization gap widens monotonically with depth (+16.6 pp jump across Layers 1–2).\n"
         "• True per-sample routing divergence correlates near-zero with downstream rescue (r = -0.05 to +0.02, p > 0.19).")
    ]

    for title, desc in audits:
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(4)
        p_a.paragraph_format.space_after = Pt(4)
        r_at = p_a.add_run(f"• {title}: ")
        r_at.bold = True
        r_at.font.color.rgb = RGBColor(0x2B, 0x5C, 0x8F)
        p_a.add_run(desc)

    # -------------------------------------------------------------
    # 4. Data Protocol, Duplicates, and Leakage
    # -------------------------------------------------------------
    add_styled_heading(doc, "4. Data Protocol, Duplicate Contamination & Leakage Analysis", 1)
    p_leak = doc.add_paragraph()
    p_leak.add_run(
        "A critical contribution of the A5-R audit is establishing the first comprehensive cryptographic audit of cross-split duplication in FER2013. "
        "Researchers have historically treated FER2013 splits as independent out-of-sample evaluations. A5-R refutes this assumption."
    )

    tbl_leak = doc.add_table(rows=4, cols=6)
    tbl_leak.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_leak = ["Split", "Total Rows", "Train Duplicates", "Leaked %", "Same Label Dup", "Conflicting Label Dup"]
    for i, h in enumerate(headers_leak):
        cell = tbl_leak.rows[0].cells[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_leak = [
        ["PublicTest (val.csv)", "3,589", "280", "7.80%", "267 (7.44%)", "13 (0.36%)"],
        ["PrivateTest (test.csv)", "3,589", "288", "8.02%", "273 (7.61%)", "15 (0.42%)"],
        ["Combined Held-Out", "7,178", "568", "7.91%", "540 (7.52%)", "28 (0.39%)"],
    ]
    for row_idx, row_vals in enumerate(data_leak):
        row = tbl_leak.rows[row_idx + 1]
        bg = "F9FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_vals):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(8)

    p_leak_disc = doc.add_paragraph()
    p_leak_disc.paragraph_format.space_before = Pt(6)
    p_leak_disc.add_run(
        "Behavioral Consequence: On exact duplicates where the label agrees, all MPG models achieve 97.0% to 100.0% accuracy (v1 achieves 273/273 on PrivateTest). "
        "Conversely, on exact duplicates where the training label contradicts the test label, accuracy collapses to 13.3%–38.5% because models memorized the training annotation. "
        "Excluding exact Train duplicates reduces nominal test accuracy by 2.1 to 2.4 percentage points across all model variants."
    )

    # -------------------------------------------------------------
    # 5. Scientific Decision Ledger & Q1 Experimental Strategy
    # -------------------------------------------------------------
    add_styled_heading(doc, "5. Scientific Decision Ledger & Q1 Experimental Strategy", 1)
    p_strat = doc.add_paragraph()
    p_strat.add_run(
        "Based on the cumulative evidence from audits A1 through A6-R2, the project establishes a binding research ledger for Q1 modeling and paper preparation:"
    )

    p_keep = doc.add_paragraph()
    r_k = p_keep.add_run("1. Confirmed Kept Mechanisms (KEEP):\n")
    r_k.bold = True
    r_k.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_keep.add_run(
        "• Dynamic Top-K Sparse Routing: Retain schedule [8, 16, 16, 16, 24]. Achieves statistical parity with dense models while reducing relational computation by >60%.\n"
        "• Model EMA: Retain decay 0.999. Essential for smooth feature geometry and validation stability.\n"
        "• Early Depth Residual Preservation: v2.3 residual scaling (res_scale=0.5 in Layers 1–2) is confirmed effective, elevating Private TTA to 70.66%.\n"
        "• Dual Readout Architecture: Retain Fusion concatenation of pixel and motif readouts."
    )

    p_dep = doc.add_paragraph()
    r_d = p_dep.add_run("2. Formally Deprioritized / Forbidden Branches (DO NOT PURSUE):\n")
    r_d.bold = True
    r_d.font.color.rgb = RGBColor(0xA9, 0x1D, 0x22) # Red
    p_dep.add_run(
        "• Graph Density Tuning / Top-K Optimization: A4, A6, and A6-R2 prove that topology density is not limiting. Do not perform grid searches over Top-K.\n"
        "• Static Geometric Graphs: A2 proved that fixed Chebyshev priors cause immediate performance degradation. Do not reintroduce static graphs.\n"
        "• Isolated Prototype / TYPE Tuning: A1 and A6-R prove that prototype assignments provide <0.5% incremental functional value over visual appearance (WHAT). Do not tune codebook size or prototype clustering losses.\n"
        "• Classifier Head Redesign: A4 and A6 prove that linear probes underperform the trained head and classifier continuation accounts for <1.2% marginal divergence. Do not modify the classifier head.\n"
        "• Automatic Dataset Relabeling: Strictly rejected. Changing benchmark ground truth introduces circular bias."
    )

    p_q1 = doc.add_paragraph()
    r_q = p_q1.add_run("3. Q1 Research Priorities (P0 & P1):\n")
    r_q.bold = True
    r_q.font.color.rgb = RGBColor(0x2B, 0x5C, 0x8F)
    p_q1.add_run(
        "• Multi-Seed Replication of v2.3 (P0): Replicate v2.3 across seeds 43 and 44 on Kaggle Tesla T4 to verify statistical robustness of the 70.66% result.\n"
        "• Paper Reporting Protocol (P0): Draft Section 4 of the paper reporting both standard FER2013 scores and duplicate-excluded scores with complete cryptographic transparency.\n"
        "• Blinded Human Review Adjudication (P1): Engage two independent FACS-trained human annotators to complete a5_human_review_form.csv without unblinding, providing definitive human perceptual calibration.\n"
        "• Resolving the 11 Common Representation Failures (P1): Perform gradient saliency on the 11 verified clean representation failure samples to determine why spatial motif pooling missed obvious expression cues."
    )

    # Save Word document
    out_docx = STAGE_DIR / "00_MASTER_SUMMARY_MPG_FER.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx} successfully!", flush=True)


if __name__ == "__main__":
    main()
